#!/usr/bin/env python3
"""
indexar_corpus.py — Extrai texto do corpus e salva em corpus_index.json.

Modo incremental: se corpus_index.json já existir, preserva entradas
anteriores e processa apenas os arquivos ainda não indexados.

Configuração obrigatória: ajuste CORPUS_ROOT antes de executar.
"""

import ast
import importlib.util
import inspect
import json
import subprocess
import sys
from datetime import datetime
from pathlib import Path

# ─── Configuração ─────────────────────────────────────────────────────────────

CORPUS_ROOT = Path(r"\\servidor\compartilhado")   # <-- AJUSTAR antes de executar
OUTPUT_FILE = Path(__file__).parent / "corpus_index.json"
DOC_COM_SCRIPT = Path(r"C:\Users\pc\extract_doc_com.py")

# Diretórios excluídos (relativos ao CORPUS_ROOT, separador normalizado para /)
EXCLUDED_DIRS: set[str] = {
    "P1/Holerites",
    "P1/Holerites Bonus DEJEM",
    "P4/endereços de ip",
}

# Extensões excluídas (minúsculas)
EXCLUDED_EXTENSIONS: set[str] = {
    ".pfx", ".cer", ".rar", ".mp4",
    ".jpg", ".jpeg", ".png", ".pptx", ".ppt",
}

# Arquivos excluídos pelo nome relativo ao CORPUS_ROOT (normalizado)
EXCLUDED_FILES: set[str] = {
    "despachadora.py",
}

# ─── Lógica de exclusão ───────────────────────────────────────────────────────

def _rel_fwd(path: Path) -> str:
    """Caminho relativo ao corpus com separadores forward-slash."""
    return path.relative_to(CORPUS_ROOT).as_posix()


def deve_excluir(path: Path) -> bool:
    rel = _rel_fwd(path)

    if path.suffix.lower() in EXCLUDED_EXTENSIONS:
        return True

    # Arquivo raiz excluído por nome
    if rel in EXCLUDED_FILES:
        return True

    # Pasta excluída: verifica se algum prefixo da hierarquia bate
    for excl in EXCLUDED_DIRS:
        if rel == excl or rel.startswith(excl + "/"):
            return True

    return False


def secao_de(path: Path) -> str:
    """Retorna o diretório de primeiro nível (seção) do arquivo."""
    partes = path.relative_to(CORPUS_ROOT).parts
    return partes[0] if len(partes) > 1 else "__raiz__"


# ─── Extrator PDF ─────────────────────────────────────────────────────────────

def _extrair_pdf(path: Path) -> tuple[str, str | None]:
    # Tentativa 1: PyMuPDF
    try:
        import fitz  # type: ignore
        partes: list[str] = []
        with fitz.open(str(path)) as doc:
            for pagina in doc:
                partes.append(pagina.get_text())
        texto = "\n".join(partes).strip()
        if texto:
            return texto, None
    except ImportError:
        pass
    except Exception:
        pass

    # Tentativa 2: pypdf
    try:
        from pypdf import PdfReader  # type: ignore
        reader = PdfReader(str(path))
        partes = []
        for pagina in reader.pages:
            t = pagina.extract_text()
            if t:
                partes.append(t)
        texto = "\n".join(partes).strip()
        if texto:
            return texto, None
    except ImportError:
        pass
    except Exception:
        pass

    return "", "pdf_imagem_sem_ocr"


# ─── Extrator DOCX ────────────────────────────────────────────────────────────

def _extrair_docx(path: Path) -> tuple[str, str | None]:
    try:
        from docx import Document  # type: ignore
        doc = Document(str(path))
        texto = "\n".join(p.text for p in doc.paragraphs).strip()
        return texto, None
    except Exception as exc:
        return "", f"docx_erro:{type(exc).__name__}"


# ─── Extrator DOC (via extract_doc_com.py) ────────────────────────────────────

# Estado do módulo COM (inicializado uma única vez)
_doc_com_state: dict = {"checked": False, "func": None, "use_subprocess": False}


def _inspecionar_doc_com_script() -> bool:
    """
    Verifica se extract_doc_com.py define funções no nível do módulo.
    Usa AST para não executar o script durante a inspeção.
    Retorna True se encontrar pelo menos uma definição de função.
    """
    try:
        src = DOC_COM_SCRIPT.read_text(encoding="utf-8", errors="replace")
        arvore = ast.parse(src)
        return any(
            isinstance(no, ast.FunctionDef)
            for no in ast.iter_child_nodes(arvore)
        )
    except Exception:
        return False


def _preparar_doc_com() -> None:
    if _doc_com_state["checked"]:
        return
    _doc_com_state["checked"] = True

    if not DOC_COM_SCRIPT.exists():
        return  # func=None, use_subprocess=False → retornará erro de script não encontrado

    tem_funcoes = _inspecionar_doc_com_script()

    if tem_funcoes:
        # Tenta importar e localizar a função extratora
        try:
            spec = importlib.util.spec_from_file_location("extract_doc_com", DOC_COM_SCRIPT)
            mod = importlib.util.module_from_spec(spec)  # type: ignore[arg-type]
            spec.loader.exec_module(mod)  # type: ignore[union-attr]

            # Busca primeira função pública com pelo menos 1 parâmetro
            for nome, obj in inspect.getmembers(mod, inspect.isfunction):
                if obj.__module__ != "extract_doc_com":
                    continue
                if nome.startswith("_"):
                    continue
                sig = inspect.signature(obj)
                if sig.parameters:
                    _doc_com_state["func"] = obj
                    return
        except Exception:
            pass

    # Se não tem funções ou import falhou → chama como subprocess
    _doc_com_state["use_subprocess"] = True


def _extrair_doc(path: Path) -> tuple[str, str | None]:
    _preparar_doc_com()

    if not DOC_COM_SCRIPT.exists():
        return "", "doc_com_script_nao_encontrado"

    func = _doc_com_state["func"]
    if func is not None:
        try:
            resultado = func(str(path))
            if isinstance(resultado, str):
                texto = resultado.strip()
                return texto, None if texto else "doc_com_vazio"
            return "", "doc_com_retorno_invalido"
        except Exception:
            return "", "doc_com_falhou"

    if _doc_com_state["use_subprocess"]:
        try:
            proc = subprocess.run(
                [sys.executable, str(DOC_COM_SCRIPT), str(path)],
                capture_output=True,
                text=True,
                timeout=60,
                encoding="utf-8",
                errors="replace",
            )
            texto = proc.stdout.strip()
            if proc.returncode == 0 and texto:
                return texto, None
            return "", "doc_com_falhou"
        except subprocess.TimeoutExpired:
            return "", "doc_com_falhou"
        except Exception:
            return "", "doc_com_falhou"

    return "", "doc_com_falhou"


# ─── Extratores de planilha ───────────────────────────────────────────────────

def _extrair_xlsx(path: Path) -> tuple[str, str | None]:
    try:
        import openpyxl  # type: ignore
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        linhas: list[str] = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                linha = "\t".join("" if c is None else str(c) for c in row)
                if linha.strip():
                    linhas.append(linha)
        wb.close()
        return "\n".join(linhas).strip(), None
    except Exception as exc:
        return "", f"xlsx_erro:{type(exc).__name__}"


def _extrair_xls(path: Path) -> tuple[str, str | None]:
    try:
        import xlrd  # type: ignore
        wb = xlrd.open_workbook(str(path))
        linhas: list[str] = []
        for sheet in wb.sheets():
            for i in range(sheet.nrows):
                linha = "\t".join(str(sheet.cell_value(i, j)) for j in range(sheet.ncols))
                if linha.strip():
                    linhas.append(linha)
        return "\n".join(linhas).strip(), None
    except Exception as exc:
        return "", f"xls_erro:{type(exc).__name__}"


# ─── Extrator de texto puro ───────────────────────────────────────────────────

_ENCODINGS_TENTATIVA = ("utf-8-sig", "utf-8", "latin-1", "cp1252")

_EXTENSOES_TEXTO: set[str] = {
    ".txt", ".md", ".csv", ".xml", ".html", ".htm",
    ".py", ".js", ".ts", ".json", ".sql",
    ".yaml", ".yml", ".ini", ".cfg", ".log",
    ".bat", ".ps1", ".sh", ".rtf",
}


def _extrair_texto(path: Path) -> tuple[str, str | None]:
    for enc in _ENCODINGS_TENTATIVA:
        try:
            return path.read_text(encoding=enc).strip(), None
        except UnicodeDecodeError:
            continue
        except Exception as exc:
            return "", f"texto_erro:{type(exc).__name__}"
    return "", "texto_encoding_falhou"


# ─── Despachante principal ────────────────────────────────────────────────────

def extrair_conteudo(path: Path) -> tuple[str, str | None]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _extrair_pdf(path)
    if ext == ".docx":
        return _extrair_docx(path)
    if ext == ".doc":
        return _extrair_doc(path)
    if ext == ".xlsx":
        return _extrair_xlsx(path)
    if ext == ".xls":
        return _extrair_xls(path)
    if ext in _EXTENSOES_TEXTO:
        return _extrair_texto(path)
    return "", f"extensao_nao_suportada:{ext}"


# ─── Gerenciamento do índice ──────────────────────────────────────────────────

def carregar_index() -> dict:
    if OUTPUT_FILE.exists():
        try:
            with OUTPUT_FILE.open(encoding="utf-8") as f:
                dados = json.load(f)
            if isinstance(dados, dict) and "entradas" in dados:
                print(f"Índice existente carregado: {len(dados['entradas'])} entradas")
                return dados
        except Exception as exc:
            print(f"Aviso: não foi possível ler {OUTPUT_FILE} ({exc}). Iniciando novo índice.")
    return {"entradas": {}, "meta": {}}


def salvar_index(dados: dict) -> None:
    with OUTPUT_FILE.open("w", encoding="utf-8") as f:
        json.dump(dados, f, ensure_ascii=False, indent=2)


# ─── Listagem de arquivos ─────────────────────────────────────────────────────

def listar_arquivos() -> list[Path]:
    resultado: list[Path] = []
    for path in sorted(CORPUS_ROOT.rglob("*")):
        if not path.is_file():
            continue
        try:
            if deve_excluir(path):
                continue
        except ValueError:
            # relative_to() pode lançar ValueError em edge cases de symlinks
            continue
        resultado.append(path)
    return resultado


# ─── Indexação principal ──────────────────────────────────────────────────────

def indexar() -> None:
    if not CORPUS_ROOT.exists():
        print(f"ERRO: CORPUS_ROOT não encontrado: {CORPUS_ROOT}", file=sys.stderr)
        sys.exit(1)

    dados = carregar_index()
    entradas: dict = dados["entradas"]

    print("Listando arquivos do corpus...")
    todos_arquivos = listar_arquivos()
    indexados = set(entradas.keys())
    novos = [a for a in todos_arquivos if str(a) not in indexados]

    print(f"Arquivos no corpus : {len(todos_arquivos)}")
    print(f"Já indexados       : {len(indexados)}")
    print(f"Novos a processar  : {len(novos)}")
    print()

    for i, path in enumerate(novos, 1):
        chave = str(path)
        sec = secao_de(path)
        rotulo = f"[{i}/{len(novos)}] {path.name[:60]:<60}  ({sec})"
        print(rotulo, end="  ", flush=True)

        try:
            texto, erro = extrair_conteudo(path)
        except Exception as exc:
            texto, erro = "", f"excecao_inesperada:{type(exc).__name__}"

        entradas[chave] = {
            "path": chave,
            "section": sec,
            "extension": path.suffix.lower(),
            "chars": len(texto),
            "text": texto,
            "error": erro,
            "indexed_at": datetime.now().isoformat(timespec="seconds"),
        }

        print(f"ERRO:{erro}" if erro else f"ok  ({len(texto):,} chars)")

        # Salva progressivamente a cada 50 arquivos para não perder progresso
        if i % 50 == 0:
            dados["meta"]["last_run"] = datetime.now().isoformat(timespec="seconds")
            salvar_index(dados)

    dados["meta"]["last_run"] = datetime.now().isoformat(timespec="seconds")
    dados["meta"]["total_files"] = len(entradas)
    salvar_index(dados)

    # ─── Resumo final ─────────────────────────────────────────────────────────
    print()
    print("=" * 65)
    print("RESUMO FINAL")
    print("=" * 65)

    total_chars = sum(e["chars"] for e in entradas.values())

    contagem_secao: dict[str, int] = {}
    contagem_erro: dict[str, int] = {}
    for e in entradas.values():
        contagem_secao[e["section"]] = contagem_secao.get(e["section"], 0) + 1
        if e["error"]:
            contagem_erro[e["error"]] = contagem_erro.get(e["error"], 0) + 1

    print(f"Total de entradas  : {len(entradas)}")
    print(f"Total de caracteres: {total_chars:,}")
    print()
    print("Por seção:")
    for sec, cnt in sorted(contagem_secao.items()):
        print(f"  {sec:<45} {cnt:>6}")
    if contagem_erro:
        print()
        print("Por tipo de erro:")
        for err, cnt in sorted(contagem_erro.items(), key=lambda x: -x[1]):
            print(f"  {err:<45} {cnt:>6}")
    else:
        print()
        print("Nenhum erro registrado.")
    print("=" * 65)
    print(f"Índice salvo em: {OUTPUT_FILE.resolve()}")


if __name__ == "__main__":
    indexar()
